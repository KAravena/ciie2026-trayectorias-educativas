#!/usr/bin/env python3
"""Reconstruct V29 native PowerPoint objects as editable Quarto/HTML/CSS/SVG.
No slide-sized screenshot is embedded. Source PPTX remains read-only.
When Quarto CLI is available: quarto render ponencia.qmd --to revealjs
Offline preview fallback: ponencia.html with local royal-runtime.js.
"""
from __future__ import annotations
import re,html,json,hashlib,csv,shutil,xml.etree.ElementTree as ET
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE,MSO_FILL_TYPE

ROOT=Path(__file__).resolve().parent
SRC=ROOT/'fuentes/CIIE2026_Aravena_Ponencia_V29_Referencias_Completas.pptx'
OLD=ROOT/'fuentes'
NOTES=json.loads((OLD/'notas_originales_V29.json').read_text())
W,H=1280,720
TITLE=['Trayectorias segmentadas de egresados técnico-profesionales hacia la educación superior chilena',
'En Chile, la transición conecta dos niveles diferenciados',
'La matrícula total no muestra cómo se distribuyen los destinos',
'El egreso es un momento de una trayectoria más amplia',
'Las trayectorias previas importan; el destino también diferencia',
'Sabemos que existen desigualdades; falta precisar esta comparación',
'Caracterizamos la continuidad y los primeros destinos según modalidad',
'Vinculamos egresos escolares con matrícula superior',
'La transición se observa mediante dos resultados',
'Cada resultado se analiza con un modelo distinto',
'Calculamos probabilidades bajo HC y TP para compararlas',
'La matrícula observada es menor entre egresados TP',
'El contraste ajustado de matrícula llega a −12,53 pp en M3',
'La distribución ajustada difiere entre primeros destinos',
'Los contrastes TP−HC varían en signo y magnitud según destino',
'Una brecha menor no implica que ambas modalidades mejoren',
'Los primeros destinos revelan una dimensión horizontal de la estratificación',
'No hay una sola brecha TP−HC: cambia según destino y cohorte',
'Referencias']
PX=lambda v:round(v/914400*96,2)

def esc(x):return html.escape(str(x),quote=True)
def color(c,default='transparent'):
 try:
  if c and c.type==MSO_COLOR_TYPE.RGB:return '#'+str(c.rgb)
  if c and c.type==MSO_COLOR_TYPE.SCHEME:return default
 except (AttributeError,ValueError,TypeError):pass
 return default

def background(sh):
 try:
  if sh.fill.type==MSO_FILL_TYPE.SOLID:return color(sh.fill.fore_color,'transparent')
 except (AttributeError,ValueError):pass
 return 'transparent'

def stroke(sh):
 try:
  c=color(sh.line.color,'none')
  if sh.line.fill.type is None:return 'none',0
  return c,PX(sh.line.width) if sh.line.width else 1
 except (AttributeError,ValueError,TypeError):return 'none',0

def xy(sh):return [PX(x) for x in (sh.left,sh.top,sh.width,sh.height)]

def font_style(run,paragraph):
 f=run.font
 family=f.name or 'Arial'
 if family.lower() in ('aptos','arial','calibri','georgia'):
  family={'aptos':'Arial, sans-serif','arial':'Arial, sans-serif','calibri':'Arial, sans-serif','georgia':"Georgia, 'Liberation Serif', serif"}[family.lower()]
 else:family=esc(family).replace('"','')+', Arial, sans-serif'
 px=round(f.size.pt*96/72,2) if f.size else 19
 clr=color(f.color,'#342D40')
 weight='700' if f.bold is True else '400'
 italic='italic' if f.italic is True else 'normal'
 under='text-decoration:underline;' if f.underline is True else ''
 return f"font-family:{family};font-size:{px}px;color:{clr};font-weight:{weight};font-style:{italic};{under}"

def text_node(sh,n):
 x,y,w,h=xy(sh)
 align={'CENTER':'center','RIGHT':'right','JUSTIFY':'justify'}.get(str(sh.text_frame.paragraphs[0].alignment).split(' ')[0],'left')
 ps=[]
 for p in sh.text_frame.paragraphs:
  al={'CENTER':'center','RIGHT':'right','JUSTIFY':'justify'}.get(str(p.alignment).split(' ')[0],align)
  runs=''.join('<span style="'+esc(font_style(r,p))+'">'+esc(r.text).replace('\n','<br>').replace('\v','<br>')+'</span>' for r in p.runs)
  if not runs: runs='&#160;'
  ps.append('<div class="ppt-para" style="text-align:'+al+'">'+runs+'</div>')
 aria=esc(sh.text.strip()[:250])
 valign='display:flex;flex-direction:column;justify-content:center;' if 'MIDDLE' in str(sh.text_frame.vertical_anchor) else ''
 return f'<div class="ppt-text" data-object="{esc(sh.name)}" role="text" aria-label="{aria}" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px;{valign}">'+''.join(ps)+'</div>'

def shape_node(sh):
 x,y,w,h=xy(sh)
 fill=background(sh);scol,sw=stroke(sh)
 if sh.shape_type==MSO_SHAPE_TYPE.LINE or (sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE and (w==0 or h==0)):
  # Open XML flips for diagonal line, straight-line paths by default.
  x1,y1=x,y; x2,y2=x+w,y+h
  try:
   if sh._element.spPr.xfrm.flipV:y1,y2=y2,y1
   if sh._element.spPr.xfrm.flipH:x1,x2=x2,x1
  except (AttributeError,TypeError):pass
  if scol=='none':scol='#DDD8E5'
  return f'<svg class="raw-line" aria-hidden="true" viewBox="0 0 1280 720"><line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{scol}" stroke-width="{sw or 1}" /></svg>'
 if sh.has_text_frame and sh.text.strip():return text_node(sh,0)
 radius=''
 try:
  if sh.auto_shape_type==MSO_AUTO_SHAPE_TYPE.OVAL:radius='border-radius:50%;'
  if sh.auto_shape_type==MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:radius='border-radius:12px;'
 except (AttributeError,ValueError):pass
 border=f'border:{sw}px solid {scol};' if scol!='none' and sw>0 else ''
 return f'<div class="ppt-rect" aria-hidden="true" data-object="{esc(sh.name)}" style="left:{x}px;top:{y}px;width:{w}px;height:{h}px;background:{fill};{radius}{border}"></div>'

# References: full original 12 V29 + Breen and Goldthorpe verified in accompanying manuscript.
ref_md=(ROOT/'fuentes/CIIE2026_Referencias_APA7_Completas_V29.md').read_text()
ref_lines=[line[2:].strip() for line in ref_md.splitlines() if line.startswith('- ')][:12]
extra='Breen, R., & Goldthorpe, J. H. (1997). Explaining educational differentials: Towards a formal rational action theory. *Rationality and Society, 9*(3), 275–305. https://doi.org/10.1177/104346397009003002'
ref_lines.insert(4,extra)
assert len(ref_lines)==13
labels=['Alvarado-Urbina (2025)','BCN (2009) · Ley 20.370','BCN (2018) · Ley 21.091','Boudon (1974)','Breen y Goldthorpe (1997)','Gerber y Cheung (2008)','Kuzmanic et al. (2022)','Kuzmanic et al. (2023)','Lucas (2001)','Mare (1980)','Peralta Rojas (2024)','Triventi (2013)','Zelada et al. (2025)']

def references_slide():
 items='\n'.join(f'<li><a href="referencias.html#ref-{i+1}">{esc(label)}</a></li>' for i,label in enumerate(labels))
 return '<div class="slide-canvas"><div class="references-one"><div class="section-kicker">REFERENCIAS · MATERIAL DE CONSULTA</div><h1>Referencias</h1><p>13 obras citadas · versión completa APA 7 disponible localmente</p><ol>'+items+'</ol><a class="ref-link" href="referencias.html">Abrir bibliografía APA 7 completa</a><div class="ref-foot">Ponencia CIIE 2026 · Fuente original V29: 12 obras; se incorporó Breen y Goldthorpe (1997), citado en notas de la lámina 05.</div></div></div>'

def bibliography_page():
 import markdown_it
 md=markdown_it.MarkdownIt('commonmark',{'html':False, 'linkify':True})
 sections=[]
 for i,ref in enumerate(ref_lines):
  # Render italic and DOI as original, always linkify URL.
  rich=md.renderInline(ref)
  # Postprocess trailing hyperlink if markdown library omitted autolink.
  rich=re.sub(r'(?<!["=])(https://doi\.org/[^\s<]+|https://www\.bcn\.cl/[^\s<]+)',r'<a href="\1">\1</a>',rich) if '<a href=' not in rich else rich
  sections.append(f'<li id="ref-{i+1}">{rich}</li>')
 css='''body{font:18px/1.65 Arial,sans-serif;color:#342d40;max-width:950px;margin:0 auto;padding:45px 32px;background:#fff}h1{font:normal 44px Georgia,serif;margin:16px 0}h2{font:normal 28px Georgia,serif;margin-top:44px}p{color:#625D69}ol{padding-left:30px}li{margin:0 0 19px;padding-left:8px;overflow-wrap:anywhere}a{color:#422168}nav a{border:1px solid #422168;border-radius:5px;padding:10px 14px;text-decoration:none}header{border-top:5px solid #422168;padding-top:20px}footer{border-top:1px solid #ddd;padding-top:16px;margin-top:35px;font-size:15px}'''
 return f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Referencias APA 7 · CIIE 2026</title><style>{css}</style></head><body><header><nav><a href="ponencia.html#/18">Volver a la lámina 19</a></nav><h1>Referencias bibliográficas completas</h1><p>Katherine Aravena Herrera · CIIE 2026 · Formato APA 7</p></header><main><h2>Obras citadas ({len(ref_lines)})</h2><ol>{''.join(sections)}</ol><h2>Procedencia de las estimaciones</h2><p>Elaboración propia a partir de registros administrativos MINEDUC–SIES y material suplementario EPAA, tablas A1–A4. No se inventan identificadores de datasets individuales.</p></main><footer>La V29 incluye 12 referencias. Esta versión agrega la obra de Breen y Goldthorpe (1997), verificada contra el manuscrito adjunto y ya citada en las notas originales de la diapositiva 05.</footer></body></html>'''

ppt=Presentation(SRC)
assert len(ppt.slides)==22
sha=hashlib.sha256(SRC.read_bytes()).hexdigest()
assert sha=='5fbc1aa1e8bd7c17d7c93d183775c95570d955e4901044e675532d2abb92ff65',sha
content=[]
for i,slide in enumerate(list(ppt.slides)[:18],start=1):
 # P47.3: la identidad del congreso ya está presente en su logotipo.
 # Se elimina SOLO la caja tipográfica redundante de la portada.
 remove=lambda sh:(i==1 and sh.name=='Text 2' and sh.has_text_frame and sh.text.strip()=='VII CONGRESO INTERDISCIPLINARIO DE INVESTIGACIÓN EN EDUCACIÓN')
 nodes=[shape_node(sh) for sh in slide.shapes if not remove(sh)]
 assert len(nodes)==len(slide.shapes)-(1 if i==1 else 0)
 content.append('<div class="slide-canvas" data-slide="'+str(i)+'">\n'+'\n'.join(nodes)+'\n</div>')
# P47.1 · Marcas institucionales de portada; recursos locales cuando estén cacheados.
# Si aún no existen archivos locales, los logos originales se cargan desde sus URLs oficiales.
LOGO_CIIE_URL='https://congresociie.cl/wp-content/uploads/2026/04/CIIE-logo-septimo.png'
LOGO_FACSO_URL='https://facso.uchile.cl/dam/jcr:16208a63-f7b7-486f-9c82-85eed075421e/logoicono150.jpg'
LOGO_CIIE='assets/logo-ciie-septimo.png' if (ROOT/'assets/logo-ciie-septimo.png').exists() else LOGO_CIIE_URL
LOGO_FACSO='assets/logo-facso-uchile.jpg' if (ROOT/'assets/logo-facso-uchile.jpg').exists() else LOGO_FACSO_URL
logos_portada=f'''<div class="cover-logos" role="group" aria-label="Logotipos del VII CIIE 2026 y de la Facultad de Ciencias Sociales de la Universidad de Chile">
  <div class="cover-logo cover-logo-ciie" role="img" aria-label="Séptimo Congreso Interdisciplinario de Investigación en Educación, CIIE 2026"><span class="logo-placeholder" aria-hidden="true">CIIE <strong>2026</strong></span><img src="{LOGO_CIIE}" alt="" loading="eager" decoding="async" onload="this.classList.add('is-loaded')" onerror="this.hidden=true" /></div>
  <div class="cover-logo cover-logo-facso" role="img" aria-label="Facultad de Ciencias Sociales, Universidad de Chile"><span class="logo-placeholder" aria-hidden="true">FACSO<br>Universidad de Chile</span><img src="{LOGO_FACSO}" alt="" loading="eager" decoding="async" onload="this.classList.add('is-loaded')" onerror="this.hidden=true" /></div>
</div>'''
content[0]=content[0].replace('<div class="slide-canvas" data-slide="1">','<div class="slide-canvas" data-slide="1">\n'+logos_portada,1)
content.append(references_slide())

# P47.4: se omite intencionalmente diapositiva 19 del V29 (sensibilidad IP).
# Quarto/Reveal.js emplea exclusivamente el CSS nativo. El reproductor de
# contingencia utiliza components.css, que nunca debe insertarse en Quarto.
# Sin title/author metadata ni comentario antes del primer ##: exactamente
# P47.4: 19 diapositivas; se excluye la lámina de sensibilidad IP (original 19).
meta='''---
lang: es
format:
  revealjs:
    slide-level: 2
    theme: [default, styles/royal-academic.scss]
    css: styles/quarto-reveal.css
    width: 1280
    height: 720
    margin: 0
    center: false
    auto-stretch: false
    slide-number: c/t
    show-slide-number: all
    transition: none
    incremental: false
    navigation-mode: linear
    hash-type: number
    controls: true
    controls-layout: edges
    controls-back-arrows: faded
    controls-tutorial: false
    progress: false
    show-notes: false
---
'''
qmd=[meta]
for i,code in enumerate(content,1):
 qmd.extend([f'\n## {TITLE[i-1]} {{.academic-slide}}\n',code,'\n::: {.notes}\n'+(NOTES.get(str(i),'Bibliografía APA 7 completa disponible en referencias.html') if i!=19 else 'Lámina de consulta. La bibliografía completa se abre desde el enlace local. Se reconcilió Breen y Goldthorpe (1997) con la fuente del manuscrito.')+'\n:::\n'])
(ROOT/'ponencia.qmd').write_text('\n'.join(qmd),encoding='utf-8')
(ROOT/'referencias.html').write_text(bibliography_page(),encoding='utf-8')
(ROOT/'assets/slide_text.json').write_text(json.dumps({str(i):{'titulo':TITLE[i-1],'notas':NOTES.get(str(i),''),'texto': '\n'.join(sh.text for sh in ppt.slides[i-1].shapes if sh.has_text_frame and sh.text.strip() and not (i==1 and sh.name=='Text 2'))} for i in range(1,19)},ensure_ascii=False,indent=2))

slides=[]
for i,code in enumerate(content,1):
 notes=esc(NOTES.get(str(i),'La bibliografía APA 7 completa puede consultarse en referencias.html.'))
 slides.append(f'<section class="academic-slide" data-slide="{i}" aria-label="Diapositiva {i}: {esc(TITLE[i-1])}">\n{code}\n<aside class="notes">{notes}</aside></section>')
# CSS and JavaScript embedded in ponencia.html make this SINGLE file portable.
# Quarto uses a different stylesheet and its official Reveal.js runtime.
css_inline=(ROOT/'styles/components.css').read_text(encoding='utf-8').replace('</style','<'+chr(92)+'/style')
js_inline=(ROOT/'assets/royal-runtime.js').read_text(encoding='utf-8').replace('</script','<'+chr(92)+'/script')
html_page=f'''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="application-name" content="CIIE 2026 · Royal Academic Editorial"><title>CIIE 2026 · Trayectorias segmentadas · Katherine Aravena Herrera</title><style>{css_inline}</style></head><body><div class="reveal"><div class="slides">{''.join(slides)}</div></div><nav class="offline-navigation" aria-label="Navegación entre diapositivas"><button class="nav-arrow nav-prev" id="prev" aria-label="Diapositiva anterior" title="Anterior (←)"><svg viewBox="0 0 24 24" width="26" height="26" aria-hidden="true"><path d="M15 4 L7 12 L15 20" /></svg></button><button class="nav-arrow nav-next" id="next" aria-label="Diapositiva siguiente" title="Siguiente (→)"><svg viewBox="0 0 24 24" width="26" height="26" aria-hidden="true"><path d="M9 4 L17 12 L9 20" /></svg></button></nav><div class="offline-slide-number" id="count" role="status" aria-live="off">1 / 19</div><script>{js_inline}</script></body></html>'''
(ROOT/'ponencia.html').write_text(html_page,encoding='utf-8')
(ROOT/'_quarto.yml').write_text('project:\n  type: default\n  output-dir: _site\n  resources:\n    - referencias.html\n    - data/*.csv\n    - assets/logo-*.png\n    - assets/logo-*.jpg\nexecute:\n  echo: false\n')
print(json.dumps({'slides':len(content),'ppt_shapes':sum(len(s.shapes) for s in list(ppt.slides)[:18]),'pptx_sha256':sha,'qmd_bytes':(ROOT/'ponencia.qmd').stat().st_size,'html_bytes':(ROOT/'ponencia.html').stat().st_size,'references':len(ref_lines)},ensure_ascii=False))
