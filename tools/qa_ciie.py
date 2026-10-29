from pathlib import Path
from collections import Counter
from difflib import SequenceMatcher
from bs4 import BeautifulSoup
from pptx import Presentation
import json,re,unicodedata,hashlib,csv,fitz
from PIL import Image,ImageOps,ImageDraw,ImageFont
P=Path(__file__).resolve().parent.parent; ROOT=P
origpdf=fitz.open(P/'fuentes/CIIE2026_Aravena_Ponencia_V29_Referencias_Completas.pdf'); newpdf=fitz.open(P/'ponencia.pdf')
ppt=Presentation(P/'fuentes/CIIE2026_Aravena_Ponencia_V29_Referencias_Completas.pptx')
soup=BeautifulSoup((P/'ponencia.html').read_text(),'html.parser');htmlslides=soup.select('section.academic-slide')
notes=json.loads((P/'fuentes/notas_originales_V29.json').read_text())
science={
 'N principal':('1.233.761',[8,16,19]),
 'HC':('793.616',[8]),'TP':('440.145',[8]),
 'HC matrícula descriptiva':('82,64',[12]),'TP matrícula descriptiva':('60,19',[12]),
 'Brecha descriptiva':('22,45',[12,13]),
 'Ajustada M3':('12,53',[13,18]),
 'IC95 inferior':('13,40',[13]),'IC95 superior':('11,67',[13]),
 'Sin matrícula':('13,50',[15]),'Universidad':('24,25',[15]),
 'IP':('6,20',[15,19]),'CFT':('4,55',[15]),
 'Cohorte 2018':('14,90',[16]),'Cohorte 2023':('9,93',[16]),
 'Sensibilidad IP':('4,46',[19]), 'Sensibilidad N':('647.747',[19])}
def norm(s):
 s=unicodedata.normalize('NFKC',s).lower().replace('−','-').replace('–','-').replace('—','-')
 return re.sub(r'\s+',' ',s).strip()
rep={'p47_estado':'entrega_20_laminas_html_y_pdf; compilacion_Quarto_pendiente','origen_sha256':hashlib.sha256((P/'fuentes/CIIE2026_Aravena_Ponencia_V29_Referencias_Completas.pptx').read_bytes()).hexdigest(),
 'laminas_original':22,'laminas_p47':len(htmlslides),'paginas_pdf_p47':len(newpdf),
 'notas_preservadas':{},'texto_preservado':{},'validacion_resultados':{},'proporciones_pdf':{}}
for i in range(1,20):
 orig_boxes=[sh.text for sh in ppt.slides[i-1].shapes if sh.has_text_frame and sh.text.strip()]
 # Excepción explícita P47.3: el rótulo del Congreso se elimina sólo en portada.
 if i==1: orig_boxes=[text for text in orig_boxes if text.strip()!='VII CONGRESO INTERDISCIPLINARIO DE INVESTIGACIÓN EN EDUCACIÓN']
 dom_boxes=['\n'.join(p.get_text(separator='\n') for p in s.select('.ppt-para')) for s in htmlslides[i-1].select('.ppt-text')]
 # Each shape text must be present; text HTML preserves runs but separates paragraph slightly
 strict=(len(orig_boxes)==len(dom_boxes)) and all(norm(a)==norm(b) for a,b in zip(orig_boxes,dom_boxes))
 rep['texto_preservado'][f'{i:02}']={'pptx_cuadros':len(orig_boxes),'html_cuadros':len(dom_boxes),'texto_de_cada_cuadro_identico_normalizado':strict,'fuente_conteo':len(ppt.slides[i-1].shapes),'html_objetos':len(htmlslides[i-1].select('.ppt-text,.ppt-rect,svg.raw-line'))}
 rep['notas_preservadas'][f'{i:02}']=norm(notes[str(i)])==norm(htmlslides[i-1].select_one('aside.notes').get_text())
 a=norm(origpdf[i-1].get_text());b=norm(newpdf[i-1].get_text())
 rep['proporciones_pdf'][f'{i:02}']=round(SequenceMatcher(None,a,b).ratio(),4)
for name,(value,which) in science.items():
 rep['validacion_resultados'][name]={str(x):value in htmlslides[x-1].get_text(' ',strip=True) for x in which}
rep['referencias']={'numero':len(soup.select('.references-one li')),'html_externo':(P/'referencias.html').exists(),'incorpora_breen_goldthorpe': 'Breen' in (P/'referencias.html').read_text()}
rep['resumen']={'conteo_objetos_pptx_1_19':sum(len(s.shapes) for s in list(ppt.slides)[:19]),'conteo_objetos_html_1_19':sum(len(s.select('.ppt-text,.ppt-rect,svg.raw-line')) for s in htmlslides[:19]),'objetos_eliminados_intencionalmente':1,
 'texto_identico_laminas':sum(x['texto_de_cada_cuadro_identico_normalizado'] for x in rep['texto_preservado'].values()),'notas_identicas':sum(rep['notas_preservadas'].values()),
 'pdf_text_similarity_media':round(sum(rep['proporciones_pdf'].values())/19,4),
 'numeros_correctos':all(all(x.values()) for x in rep['validacion_resultados'].values())}
(P/'docs/qa_data.json').write_text(json.dumps(rep,ensure_ascii=False,indent=2))
print(json.dumps(rep['resumen'],ensure_ascii=False))
print('numbers',rep['validacion_resultados'])
# Side-by-side comparison 1, 5, 13, 16, 18, 19
indices=[1,5,13,16,18,19]
canvas=Image.new('RGB',(1600,3*472),'#eeeaf1');draw=ImageDraw.Draw(canvas)
font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',14)
for ix,n in enumerate(indices):
 y=ix//2*472; x=(ix%2)*800
 orig=origpdf[n-1].get_pixmap(matrix=fitz.Matrix(0.8,0.8),alpha=False)
 new=newpdf[n-1].get_pixmap(matrix=fitz.Matrix(0.8,0.8),alpha=False)
 a=Image.frombytes('RGB',[orig.width,orig.height],orig.samples)
 b=Image.frombytes('RGB',[new.width,new.height],new.samples)
 a.thumbnail((772,209));b.thumbnail((772,209))
 draw.text((x+14,y+7),f'Lámina {n:02d} · original PowerPoint',font=font,fill='#422168')
 canvas.paste(a,(x+14,y+30));draw.text((x+14,y+236),'Reconstrucción HTML/SVG → PDF',font=font,fill='#422168')
 canvas.paste(b,(x+14,y+258))
canvas.save(P/'capturas/comparacion_6_laminas_P47.png')
# gallery of entire slide set, PDF derived
thumb_w,thumb_h=370,208;gap=20;hdr=26
cols=4;rows=5
sheet=Image.new('RGB',(cols*(thumb_w+gap)+gap,rows*(thumb_h+hdr+gap)+gap),'#edeaf0');pen=ImageDraw.Draw(sheet)
for k in range(20):
 x=gap+(k%cols)*(thumb_w+gap);y=gap+(k//cols)*(thumb_h+hdr+gap)
 img=Image.open(P/'capturas'/f'lamina_{k+1:02}.png').convert('RGB');img.thumbnail((thumb_w,thumb_h))
 sheet.paste(img,(x,y+hdr));pen.text((x,y+2),f'{k+1:02d} · {"Referencias" if k==19 else "Presentación" if k<18 else "Respaldo"}',font=font,fill='#422168')
sheet.save(P/'capturas/contacto_20_laminas.png')
